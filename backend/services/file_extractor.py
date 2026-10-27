"""
PrecurAi - Multi-Format Safety Report File Extraction Service
Handles text and page extraction from PDF, DOCX, and TXT safety reports.
Validates file format, size limits, and extractable character thresholds.
"""

import io
import os
import re
import logging
from typing import Dict, List, Any, Optional

logger = logging.getLogger("precurai.file_extractor")

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class FileExtractionResult:
    def __init__(
        self,
        filename: str,
        file_type: str,
        file_size_bytes: int,
        full_text: str,
        pages: List[Dict[str, Any]],
        is_scanned: bool = False,
        scan_warning: Optional[str] = None
    ):
        self.filename = filename
        self.file_type = file_type
        self.file_size_bytes = file_size_bytes
        self.full_text = full_text
        self.pages = pages  # List of {"page": int, "text": str}
        self.is_scanned = is_scanned
        self.scan_warning = scan_warning

    def to_dict(self) -> Dict[str, Any]:
        return {
            "filename": self.filename,
            "file_type": self.file_type,
            "file_size_bytes": self.file_size_bytes,
            "total_pages": len(self.pages),
            "full_text": self.full_text,
            "pages": self.pages,
            "is_scanned": self.is_scanned,
            "scan_warning": self.scan_warning
        }


def extract_file_content(file_bytes: bytes, filename: str) -> FileExtractionResult:
    """
    Validates and extracts unstructured safety text from uploaded PDF, DOCX, or TXT files.
    """
    if not file_bytes:
        raise ValueError("Uploaded file is empty (0 bytes). Please upload a valid document.")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise ValueError(f"File size ({len(file_bytes) / (1024*1024):.1f}MB) exceeds allowable 10MB limit.")

    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file format '{ext}'. PrecurAi MVP supports .pdf, .docx, and .txt documents.")

    if ext == ".pdf":
        return _extract_pdf(file_bytes, filename)
    elif ext == ".docx":
        return _extract_docx(file_bytes, filename)
    elif ext == ".txt":
        return _extract_txt(file_bytes, filename)
    else:
        raise ValueError(f"Unsupported file format '{ext}'.")


def _extract_pdf(file_bytes: bytes, filename: str) -> FileExtractionResult:
    import fitz  # PyMuPDF

    pages: List[Dict[str, Any]] = []
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            text = page.get_text("text") or ""
            text = text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n").strip()
            pages.append({"page": page_idx + 1, "text": text})
        doc.close()
    except Exception as e:
        logger.warning(f"PyMuPDF stream extraction failed for {filename}: {e}")
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    text = text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n").strip()
                    pages.append({"page": page_idx + 1, "text": text})
        except Exception as e2:
            raise ValueError(f"Failed to parse PDF document '{filename}': {e2}")

    total_chars = 0
    for p in pages:
        if p["text"]:
            total_chars += len(re.sub(r"\s+", "", p["text"]))

    full_text = "\n\n".join([f"--- PAGE {p['page']} ---\n{p['text']}" if len(pages) > 1 else p["text"] for p in pages if p["text"]])

    if not full_text.strip() or total_chars < 15:
        raise ValueError(f"PDF document '{filename}' is unreadable or contains no extractable text. If it is a scanned image, OCR processing is required.")

    is_scanned = (total_chars / max(len(pages), 1)) < 15
    scan_warning = "Document appears to be scanned or image-based." if is_scanned else None

    return FileExtractionResult(
        filename=filename,
        file_type="pdf",
        file_size_bytes=len(file_bytes),
        full_text=full_text,
        pages=pages,
        is_scanned=is_scanned,
        scan_warning=scan_warning
    )


def _extract_docx(file_bytes: bytes, filename: str) -> FileExtractionResult:
    import docx

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_str:
                    table_texts.append(row_str)

        all_text_blocks = paragraphs + table_texts
        full_text = "\n\n".join(all_text_blocks)

        if not full_text.strip() or len(re.sub(r"\s+", "", full_text)) < 15:
            raise ValueError(f"DOCX document '{filename}' contains no extractable text.")

        pages = [{"page": 1, "text": full_text}]
        return FileExtractionResult(
            filename=filename,
            file_type="docx",
            file_size_bytes=len(file_bytes),
            full_text=full_text,
            pages=pages
        )
    except ValueError as ve:
        raise ve
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document '{filename}': {e}")


def _extract_txt(file_bytes: bytes, filename: str) -> FileExtractionResult:
    full_text = ""
    for encoding in ["utf-8", "latin-1", "cp1252"]:
        try:
            full_text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if not full_text.strip() or len(re.sub(r"\s+", "", full_text)) < 15:
        raise ValueError(f"TXT document '{filename}' is empty or contains no readable text.")

    pages = [{"page": 1, "text": full_text.strip()}]
    return FileExtractionResult(
        filename=filename,
        file_type="txt",
        file_size_bytes=len(file_bytes),
        full_text=full_text.strip(),
        pages=pages
    )
